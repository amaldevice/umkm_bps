"""
Streamlit App - EDA Usaha
=========================
Aplikasi web untuk melakukan EDA pada data usaha.

Jalankan dengan:
    streamlit run streamlit_app.py

Author: Generated for data processing automation
"""

import streamlit as st
import pandas as pd
import re
from typing import Optional, List, Dict
from io import BytesIO
import warnings
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
warnings.filterwarnings('ignore')


# =============================================================================
# PAGE CONFIG
# =============================================================================
st.set_page_config(
    page_title="EDA Usaha",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)


# =============================================================================
# HELPER FUNCTIONS
# =============================================================================

@st.cache_data
def load_data(uploaded_file):
    """Load data dari uploaded file (CSV atau Excel)."""
    file_name = uploaded_file.name.lower()
    
    if file_name.endswith(('.xlsx', '.xls')):
        df = pd.read_excel(uploaded_file)
    else:
        # Coba berbagai encoding untuk CSV
        try:
            df = pd.read_csv(uploaded_file, encoding='utf-8')
        except UnicodeDecodeError:
            try:
                uploaded_file.seek(0)
                df = pd.read_csv(uploaded_file, encoding='latin-1')
            except UnicodeDecodeError:
                uploaded_file.seek(0)
                df = pd.read_csv(uploaded_file, encoding='cp1252')
    
    return df


def clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Bersihkan nama kolom."""
    df = df.copy()
    df.columns = df.columns.str.split('\n').str[0]
    return df


def format_kode_wilayah(df: pd.DataFrame) -> pd.DataFrame:
    """Format kode wilayah dengan zero-padding."""
    df = df.copy()
    
    if 'kdprov' in df.columns:
        df['kdprov'] = df['kdprov'].astype("Int64").astype(str).str.zfill(2)
    if 'kdkab' in df.columns:
        df['kdkab'] = df['kdkab'].astype("Int64").astype(str).str.zfill(2)
    if 'kdkec' in df.columns:
        df['kdkec'] = df['kdkec'].astype("Int64").astype(str).str.zfill(3)
    if 'kddesa' in df.columns:
        df['kddesa'] = df['kddesa'].astype("Int64").astype(str).str.zfill(3)
    
    return df


def format_wilayah_lookup(df: pd.DataFrame) -> pd.DataFrame:
    """Format kode wilayah di lookup table."""
    df = df.copy()
    
    df['kdprov'] = df['kdprov'].astype("Int64").astype(str).str.zfill(2)    
    df['kdkab'] = df['kdkab'].astype("Int64").astype(str).str.zfill(2)
    df['kdkec'] = df['kdkec'].astype("Int64").astype(str).str.zfill(3)
    df['kddesa'] = df['kddesa'].astype("Int64").astype(str).str.zfill(3)
    
    return df


def create_alamat_new(df_wilayah: pd.DataFrame, kota_codes: List[str]) -> pd.DataFrame:
    """Buat kolom alamat_new dari data wilayah (untuk lookup table)."""
    df = df_wilayah.copy()
    
    # Pastikan kolom nama bertipe string
    for col in ["nmdesa", "nmkec", "nmkab", "nmprov"]:
        if col in df.columns:
            df[col] = df[col].fillna('').astype(str).str.strip()
    
    def build_alamat(row):
        parts = []
        
        if row.get('nmdesa', '') and row['nmdesa'] not in ['', 'nan', 'None', '<NA>']:
            parts.append(f"Desa {row['nmdesa']}")
        
        if row.get('nmkec', '') and row['nmkec'] not in ['', 'nan', 'None', '<NA>']:
            parts.append(f"Kecamatan {row['nmkec']}")
        
        if row.get('nmkab', '') and row['nmkab'] not in ['', 'nan', 'None', '<NA>']:
            kdkab = str(row.get('kdkab', ''))
            prefix = "Kota" if kdkab in kota_codes else "Kabupaten"
            parts.append(f"{prefix} {row['nmkab']}")
        
        if row.get('nmprov', '') and row['nmprov'] not in ['', 'nan', 'None', '<NA>']:
            parts.append(f"Provinsi {row['nmprov']}")
        
        return ', '.join(parts) if parts else None
    
    df['alamat_new'] = df.apply(build_alamat, axis=1)
    return df


def create_fallback_alamat(df: pd.DataFrame, df_wilayah: pd.DataFrame, kota_codes: List[str], alamat_col: str = 'alamat') -> pd.DataFrame:
    """Buat alamat fallback untuk rows yang tidak match. Isi kolom alamat yang sudah ada jika kosong."""
    df = df.copy()
    
    # Pastikan kolom alamat ada
    if alamat_col not in df.columns:
        df[alamat_col] = None
    
    # Hanya isi yang kosong/null/NaN
    mask_empty = df[alamat_col].isna()
    if mask_empty.any():
        # Cek juga yang string kosong atau 'nan'
        mask_str_empty = ~df[alamat_col].isna() & (df[alamat_col].astype(str).str.strip().isin(['', 'nan', 'None', '<NA>']))
        mask_empty = mask_empty | mask_str_empty
    
    if not mask_empty.any():
        return df
    
    # Lookup tables
    lookup_kec = df_wilayah[['kdprov', 'kdkab', 'kdkec', 'nmkec', 'nmkab', 'nmprov']].drop_duplicates(
        subset=['kdprov', 'kdkab', 'kdkec']
    )
    lookup_kab = df_wilayah[['kdprov', 'kdkab', 'nmkab', 'nmprov']].drop_duplicates(
        subset=['kdprov', 'kdkab']
    )
    lookup_prov = df_wilayah[['kdprov', 'nmprov']].drop_duplicates(subset=['kdprov'])
    
    def build_fallback(row):
        parts = []
        kdprov = str(row.get('kdprov', ''))
        kdkab = str(row.get('kdkab', ''))
        kdkec = str(row.get('kdkec', ''))
        
        # Coba match level kecamatan
        match_kec = lookup_kec[
            (lookup_kec['kdprov'] == kdprov) & 
            (lookup_kec['kdkab'] == kdkab) & 
            (lookup_kec['kdkec'] == kdkec)
        ]
        
        if not match_kec.empty:
            r = match_kec.iloc[0]
            if r.get('nmkec', '') not in ['', 'nan', 'None']:
                parts.append(f"Kecamatan {r['nmkec']}")
            if r.get('nmkab', '') not in ['', 'nan', 'None']:
                prefix = "Kota" if kdkab in kota_codes else "Kabupaten"
                parts.append(f"{prefix} {r['nmkab']}")
            if r.get('nmprov', '') not in ['', 'nan', 'None']:
                parts.append(f"Provinsi {r['nmprov']}")
            if parts:
                return ', '.join(parts)
        
        # Coba match level kabupaten
        match_kab = lookup_kab[
            (lookup_kab['kdprov'] == kdprov) & 
            (lookup_kab['kdkab'] == kdkab)
        ]
        
        if not match_kab.empty:
            r = match_kab.iloc[0]
            if r.get('nmkab', '') not in ['', 'nan', 'None']:
                prefix = "Kota" if kdkab in kota_codes else "Kabupaten"
                parts.append(f"{prefix} {r['nmkab']}")
            if r.get('nmprov', '') not in ['', 'nan', 'None']:
                parts.append(f"Provinsi {r['nmprov']}")
            if parts:
                return ', '.join(parts)
        
        # Coba match level provinsi
        match_prov = lookup_prov[lookup_prov['kdprov'] == kdprov]
        if not match_prov.empty:
            nmprov = match_prov.iloc[0].get('nmprov', '')
            if nmprov not in ['', 'nan', 'None']:
                return f"Provinsi {nmprov}"
        
        # Fallback ke kode
        codes = []
        if kdprov: codes.append(f"Prov:{kdprov}")
        if kdkab: codes.append(f"Kab:{kdkab}")
        if kdkec: codes.append(f"Kec:{kdkec}")
        kddesa = str(row.get('kddesa', ''))
        if kddesa: codes.append(f"Desa:{kddesa}")
        
        return f"[Kode: {', '.join(codes)}]" if codes else None
    
    # Isi kolom alamat yang kosong dengan fallback
    df.loc[mask_empty, alamat_col] = df.loc[mask_empty].apply(build_fallback, axis=1)
    return df


def merge_with_wilayah(df: pd.DataFrame, df_wilayah: pd.DataFrame, kota_codes: List[str], alamat_col: str = 'alamat') -> pd.DataFrame:
    """Merge data usaha dengan data wilayah. Isi kolom alamat yang sudah ada jika kosong."""
    df = df.copy()
    
    # Pastikan kolom alamat ada
    if alamat_col not in df.columns:
        df[alamat_col] = None
    
    df_wilayah = format_wilayah_lookup(df_wilayah)
    df_wilayah = create_alamat_new(df_wilayah, kota_codes)
    
    on_columns = ["kdprov", "kdkab", "kdkec", "kddesa"]
    lookup = df_wilayah[on_columns + ["alamat_new"]].drop_duplicates(subset=on_columns)
    
    # Merge untuk mendapatkan alamat_new dari lookup
    df = df.merge(lookup, on=on_columns, how='left')
    
    # Isi kolom alamat yang kosong dengan alamat_new
    mask_empty = df[alamat_col].isna()
    if mask_empty.any():
        # Cek juga yang string kosong atau 'nan'
        mask_str_empty = ~df[alamat_col].isna() & (df[alamat_col].astype(str).str.strip().isin(['', 'nan', 'None', '<NA>']))
        mask_empty = mask_empty | mask_str_empty
    
    df.loc[mask_empty & df['alamat_new'].notna(), alamat_col] = df.loc[mask_empty & df['alamat_new'].notna(), 'alamat_new']
    
    # Hapus kolom alamat_new yang temporary
    if 'alamat_new' in df.columns:
        df = df.drop(columns=['alamat_new'])
    
    # Fallback untuk yang masih kosong
    df = create_fallback_alamat(df, df_wilayah, kota_codes, alamat_col)
    
    return df


def get_bentuk_badan_hukum_patterns() -> dict:
    """Definisi pattern regex untuk klasifikasi bentuk badan hukum."""
    patterns = {
        1: r'(?i)\b(PT\.?|P\.T\.?|PERSEROAN|PERSERO)\b',
        2: r'(?i)\b(YAYASAN|YYS\.?)\b',
        3: r'(?i)\b(KOPERASI|KOP\.?|KOPKAR|KSU|KSP|KOPWAN|KOPDIT|KOSPIN|KUD)\b',
        4: r'(?i)\b(DANA\s*PENSIUN|DAPEN|DP\.?)\b',
        5: r'(?i)\b(PERUM|PERUMDA|PERUMNAS)\b',
        6: r'(?i)\b(BUMDES|BUM\s*DESA|BUMN?\s*DES)\b',
        7: r'(?i)\b(CV\.?|C\.V\.?)\b',
        8: r'(?i)\b(FIRMA|FA\.?)\b',
        9: r'(?i)\b(PERSEKUTUAN|MAATSCHAP|SEKUTU)\b',
        10: r'(?i)\b(PERWAKILAN|REPRESENTATIVE\s*OFFICE|REP\.?\s*OFFICE)\b',
        11: r'(?i)\b(LLC|LTD\.?|LIMITED|INC\.?|INCORPORATED|CORP\.?|CORPORATION|PTE\.?\s*LTD\.?|SDN\.?\s*BHD\.?|GMBH|AG\b|S\.?A\.?\b|BV\.?|NV\.?)\b',
        12: r'(?i)\b(UD\.?|U\.D\.?|PD\.?|P\.D\.?|TOKO|WARUNG|KIOS|KEDAI|DEPOT|SALON|BENGKEL|LAUNDRY|FOTOCOPY|FOTO\s*COPY|PERCETAKAN|SABLON|KONTER|COUNTER|RENTAL|SEWA|JAHIT|TAILOR|CATERING|KATERING)\b',
    }
    return patterns


def classify_bentuk_badan_hukum(df: pd.DataFrame, nama_col: str = 'nama_usaha') -> pd.DataFrame:
    """Klasifikasi bentuk badan hukum berdasarkan nama usaha. Overwrite kolom bentuk_badan_hukum_usaha."""
    df = df.copy()
    patterns = get_bentuk_badan_hukum_patterns()
    
    # Pastikan kolom bentuk_badan_hukum_usaha ada
    if 'bentuk_badan_hukum_usaha' not in df.columns:
        df['bentuk_badan_hukum_usaha'] = None
    
    def classify_single(nama):
        if pd.isna(nama) or str(nama).strip() == '':
            return None
        nama_str = str(nama).upper()
        for kode, pattern in patterns.items():
            if re.search(pattern, nama_str):
                return kode
        return None
    
    # Overwrite kolom bentuk_badan_hukum_usaha
    df['bentuk_badan_hukum_usaha'] = df[nama_col].apply(classify_single)
    return df


def get_kode_names() -> dict:
    """Mapping kode ke nama bentuk badan hukum."""
    return {
        1: 'Perseroan (PT)',
        2: 'Yayasan',
        3: 'Koperasi',
        4: 'Dana Pensiun',
        5: 'Perum/Perumda',
        6: 'BUM Desa',
        7: 'CV',
        8: 'Firma',
        9: 'Persekutuan Perdata',
        10: 'Kantor Perwakilan LN',
        11: 'Badan Usaha LN',
        12: 'Usaha Perseorangan',
        99: 'Lainnya'
    }


def organize_columns_by_category(df: pd.DataFrame) -> tuple:
    """Organisir kolom berdasarkan kategori yang ditentukan.
    
    Returns:
        tuple: (df_ordered, ordered_columns, category_ranges)
    """
    df = df.copy()
    
    # Mapping kolom ke kategori berdasarkan pola nama (case-insensitive)
    categories = {
        'IDENTITAS USAHA/PERUSAHAAN': [
            'idsbr', 'nama_usaha', 'nama_komersial_usaha', 'alamat', 'nama_sls',
            'kodepos', 'nomor_te', 'nomor_wemail', 'website', 'latitude', 'longitude'
        ],
        'KEBERADAAN USAHA/PERUSAHAAN': [
            'keberadaan_usaha', 'keberadaan_usaha/perusahaan', 'idsbr_master',
            'kdprov_pindah', 'kdkab_pindah', 'kdkec_pindah', 'kddesa_pindah'
        ],
        'WILAYAH USAHA/PERUSAHAAN': [
            'kdprov', 'kdkab', 'kdkec', 'kddesa', 'nmprov', 'nmkab', 'nmkec', 'nmdesa'
        ],
        'KARAKTERISTIK USAHA/PERUSAHAAN': [
            'bentuk_badan_hukum_usaha', 'status_usaha', 'jenis_usaha',
            'kategori_usaha', 'skala_usaha', 'modal_usaha'
        ],
        'KEGIATAN USAHA/PERUSAHAAN': [
            'kode_kbli', 'nama_kbli', 'kegiatan_utama', 'kegiatan_tambahan'
        ],
        'LAIN LAIN': []
    }
    
    # Kumpulkan semua kolom yang sudah di-mapping
    mapped_columns = set()
    for cat_cols in categories.values():
        mapped_columns.update(cat_cols)
    
    # Organisir kolom sesuai urutan kategori
    ordered_columns = []
    category_ranges = []  # Menyimpan (category_name, start_idx, end_idx)
    
    # Helper function untuk mendapatkan kolom dari kategori
    def get_category_columns(category_name, max_cols):
        cat_cols = []
        # Pertama, cari kolom yang exact match dengan list
        for col in categories[category_name]:
            if col in df.columns and col not in ordered_columns and len(cat_cols) < max_cols:
                cat_cols.append(col)
        # Jika masih ada slot, cari kolom yang mirip (case-insensitive, partial match)
        if len(cat_cols) < max_cols:
            for col in df.columns:
                if col in ordered_columns:
                    continue
                col_lower = col.lower()
                for pattern in categories[category_name]:
                    pattern_lower = pattern.lower()
                    if (pattern_lower in col_lower or col_lower in pattern_lower) and len(cat_cols) < max_cols:
                        cat_cols.append(col)
                        break
        return cat_cols
    
    # 1. IDENTITAS USAHA/PERUSAHAAN (A-L, maks 12 kolom)
    identitas_cols = get_category_columns('IDENTITAS USAHA/PERUSAHAAN', 12)
    start_idx = len(ordered_columns)
    ordered_columns.extend(identitas_cols)
    if identitas_cols:
        category_ranges.append(('IDENTITAS USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    # 2. KEBERADAAN USAHA/PERUSAHAAN (M-P, maks 4 kolom)
    keberadaan_cols = get_category_columns('KEBERADAAN USAHA/PERUSAHAAN', 4)
    start_idx = len(ordered_columns)
    ordered_columns.extend(keberadaan_cols)
    if keberadaan_cols:
        category_ranges.append(('KEBERADAAN USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    # 3. WILAYAH USAHA/PERUSAHAAN (Q-T, maks 4 kolom)
    wilayah_cols = get_category_columns('WILAYAH USAHA/PERUSAHAAN', 4)
    start_idx = len(ordered_columns)
    ordered_columns.extend(wilayah_cols)
    if wilayah_cols:
        category_ranges.append(('WILAYAH USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    # 4. KARAKTERISTIK USAHA/PERUSAHAAN (U-Z, maks 6 kolom)
    karakteristik_cols = get_category_columns('KARAKTERISTIK USAHA/PERUSAHAAN', 6)
    start_idx = len(ordered_columns)
    ordered_columns.extend(karakteristik_cols)
    if karakteristik_cols:
        category_ranges.append(('KARAKTERISTIK USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    # 5. KEGIATAN USAHA/PERUSAHAAN (AA-AD, maks 4 kolom)
    kegiatan_cols = get_category_columns('KEGIATAN USAHA/PERUSAHAAN', 4)
    start_idx = len(ordered_columns)
    ordered_columns.extend(kegiatan_cols)
    if kegiatan_cols:
        category_ranges.append(('KEGIATAN USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    # 6. Kolom lainnya yang belum di-mapping (LAIN LAIN, maks 2 kolom)
    remaining_cols = [col for col in df.columns if col not in ordered_columns][:2]
    start_idx = len(ordered_columns)
    ordered_columns.extend(remaining_cols)
    if remaining_cols:
        category_ranges.append(('LAIN LAIN', start_idx, len(ordered_columns) - 1))
    
    # Tambahkan kolom yang belum masuk (jika ada)
    all_remaining = [col for col in df.columns if col not in ordered_columns]
    ordered_columns.extend(all_remaining)
    
    # Reorder dataframe
    df_ordered = df[ordered_columns]
    
    return df_ordered, ordered_columns, category_ranges


def to_excel_bytes(df: pd.DataFrame) -> bytes:
    """Convert DataFrame ke Excel bytes dengan format terorganisir.
    
    - Urutan kolom DIJAGA sama persis dengan DataFrame input (seperti CSV).
    - Header kategori hanya menandai blok kolom:
      A-L: IDENTITAS, M-P: KEBERADAAN, Q-T: WILAYAH,
      U-Z: KARAKTERISTIK, AA-AD: KEGIATAN, AE-AF: LAIN LAIN.
    """
    output = BytesIO()
    
    # Gunakan urutan kolom apa adanya dari DataFrame
    df_ordered = df.copy()
    ordered_columns = list(df_ordered.columns)
    total_cols = len(ordered_columns)
    
    # Bangun category_ranges berdasarkan posisi kolom (bukan nama)
    # Segment: (nama_kategori, jumlah_maks_kolom)
    segments = [
        ('IDENTITAS USAHA/PERUSAHAAN', 12),   # A-L
        ('KEBERADAAN USAHA/PERUSAHAAN', 4),   # M-P
        ('WILAYAH USAHA/PERUSAHAAN', 4),      # Q-T
        ('KARAKTERISTIK USAHA/PERUSAHAAN', 6),# U-Z
        ('KEGIATAN USAHA/PERUSAHAAN', 4),     # AA-AD
        ('LAIN LAIN', 2),                     # AE-AF
    ]
    category_ranges = []
    current_idx = 0
    for cat_name, width in segments:
        if current_idx >= total_cols:
            break
        start_idx = current_idx
        end_idx = min(current_idx + width - 1, total_cols - 1)
        category_ranges.append((cat_name, start_idx, end_idx))
        current_idx = end_idx + 1
    
    # Tulis ke Excel
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_ordered.to_excel(writer, index=False, sheet_name='Sheet1', startrow=1)
        
        # Ambil workbook dan worksheet
        workbook = writer.book
        worksheet = writer.sheets['Sheet1']
        
        # Buat header row untuk kategori (row 1)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        border_style = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Tulis kategori header di row 1 berdasarkan category_ranges
        for cat_name, start_idx, end_idx in category_ranges:
            start_letter = get_column_letter(start_idx + 1)  # +1 karena Excel 1-indexed
            end_letter = get_column_letter(end_idx + 1)
            
            # Merge cells untuk header kategori (jika lebih dari 1 kolom)
            if start_idx < end_idx:
                worksheet.merge_cells(f'{start_letter}1:{end_letter}1')
            
            # Set nilai dan style
            cell = worksheet[f'{start_letter}1']
            cell.value = cat_name
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border_style
            
            # Jika hanya 1 kolom, tetap set border pada cell tersebut
            if start_idx == end_idx:
                cell.border = border_style
        
        # Style untuk header kolom (row 2)
        header_row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        header_row_font = Font(bold=True, size=10)
        
        for col_idx, col_name in enumerate(ordered_columns, 1):
            col_letter = get_column_letter(col_idx)
            cell = worksheet[f'{col_letter}2']
            cell.fill = header_row_fill
            cell.font = header_row_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_style
        
        # Set row height untuk header kategori
        worksheet.row_dimensions[1].height = 30
        worksheet.row_dimensions[2].height = 20
        
        # Auto-adjust column widths
        for col_idx, col_name in enumerate(ordered_columns, 1):
            col_letter = get_column_letter(col_idx)
            max_length = max(
                len(str(col_name)),
                df_ordered[col_name].astype(str).apply(len).max() if len(df_ordered) > 0 else 0
            )
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[col_letter].width = adjusted_width
        
        # Freeze panes pada row 2 (header kolom)
        worksheet.freeze_panes = 'A3'
    
    output.seek(0)
    return output.getvalue()


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    """Convert DataFrame ke CSV bytes."""
    return df.to_csv(index=False).encode('utf-8')


# =============================================================================
# STREAMLIT UI
# =============================================================================

def main():
    st.title("📊 EDA Usaha")
    st.markdown("Aplikasi untuk melakukan Exploratory Data Analysis pada data usaha")
    
    # Sidebar
    with st.sidebar:
        st.header("⚙️ Pengaturan")
        
        # Upload files
        st.subheader("📁 Upload Data")
        
        usaha_file = st.file_uploader(
            "Data Usaha (CSV/Excel)",
            type=['csv', 'xlsx', 'xls'],
            help="Upload file data usaha"
        )
        
        wilayah_file = st.file_uploader(
            "Data Wilayah (Excel) - Opsional",
            type=['xlsx', 'xls'],
            help="Upload file lookup wilayah untuk membuat alamat_new"
        )
        
        st.divider()
        
        # Kota codes
        st.subheader("🏙️ Kode Kota")
        kota_codes_input = st.text_input(
            "Kode Kabupaten yang merupakan Kota",
            value="71",
            help="Pisahkan dengan koma, contoh: 71,72,73"
        )
        kota_codes = [k.strip() for k in kota_codes_input.split(',') if k.strip()]
        
        st.divider()
        
        # Output format
        st.subheader("📤 Format Output")
        output_format = st.selectbox(
            "Pilih format output",
            options=['Excel (.xlsx)', 'CSV (.csv)'],
            index=0
        )
        
        output_filename = st.text_input(
            "Nama file output",
            value="hasil_eda"
        )
    
    # Main content
    if usaha_file is not None:
        # Load data
        with st.spinner("Loading data usaha..."):
            df = load_data(usaha_file)
            df_original = df.copy()
        
        st.success(f"✅ Data loaded: {df.shape[0]:,} rows × {df.shape[1]} columns")
        
        # Tabs
        tab1, tab2, tab3, tab4, tab5 = st.tabs([
            "📋 Data Preview", 
            "🔧 Processing", 
            "📊 Klasifikasi Badan Hukum",
            "📈 Statistik",
            "💾 Download"
        ])
        
        # Tab 1: Data Preview
        with tab1:
            st.subheader("Data Original")
            st.dataframe(df_original.head(100), use_container_width=True)
            
            col1, col2, col3 = st.columns(3)
            with col1:
                st.metric("Total Rows", f"{len(df_original):,}")
            with col2:
                st.metric("Total Columns", len(df_original.columns))
            with col3:
                st.metric("Duplicated Rows", f"{df_original.duplicated().sum():,}")
        
        # Tab 2: Processing
        with tab2:
            st.subheader("🔧 Proses Data")
            
            # Checkboxes for processing steps
            col1, col2 = st.columns(2)
            
            with col1:
                do_clean_columns = st.checkbox("Clean column names", value=True)
                do_format_kode = st.checkbox("Format kode wilayah", value=True)
            
            with col2:
                do_merge_wilayah = st.checkbox(
                    "Merge dengan data wilayah", 
                    value=wilayah_file is not None,
                    disabled=wilayah_file is None
                )
                do_classify = st.checkbox("Klasifikasi bentuk badan hukum", value=True)
            
            if st.button("🚀 Proses Data", type="primary"):
                progress = st.progress(0)
                status = st.empty()
                
                df_processed = df.copy()
                
                # Step 1: Clean columns
                if do_clean_columns:
                    status.text("Cleaning column names...")
                    df_processed = clean_columns(df_processed)
                    progress.progress(20)
                
                # Step 2: Format kode wilayah
                if do_format_kode:
                    status.text("Formatting kode wilayah...")
                    df_processed = format_kode_wilayah(df_processed)
                    progress.progress(40)
                
                # Step 3: Merge dengan wilayah
                if do_merge_wilayah and wilayah_file is not None:
                    status.text("Merging dengan data wilayah...")
                    df_wilayah = pd.read_excel(wilayah_file)
                    # Deteksi kolom alamat (bisa alamat, alamat_usaha, atau variasi lainnya)
                    alamat_cols = [col for col in df_processed.columns if 'alamat' in col.lower()]
                    alamat_col = alamat_cols[0] if alamat_cols else 'alamat'
                    df_processed = merge_with_wilayah(df_processed, df_wilayah, kota_codes, alamat_col)
                    progress.progress(60)
                
                # Step 4: Klasifikasi
                if do_classify:
                    status.text("Classifying bentuk badan hukum...")
                    nama_col = 'nama_usaha' if 'nama_usaha' in df_processed.columns else df_processed.columns[1]
                    df_processed = classify_bentuk_badan_hukum(df_processed, nama_col)
                    progress.progress(80)
                
                progress.progress(100)
                status.text("✅ Processing completed!")
                
                # Save to session state
                st.session_state['df_processed'] = df_processed
                
                st.success(f"✅ Data processed: {df_processed.shape[0]:,} rows × {df_processed.shape[1]} columns")
                st.dataframe(df_processed.head(100), use_container_width=True)
        
        # Tab 3: Klasifikasi Badan Hukum
        with tab3:
            st.subheader("📊 Klasifikasi Bentuk Badan Hukum")
            
            if 'df_processed' in st.session_state and 'bentuk_badan_hukum_usaha' in st.session_state['df_processed'].columns:
                df_proc = st.session_state['df_processed']
                
                # Summary
                kode_names = get_kode_names()
                value_counts = df_proc['bentuk_badan_hukum_usaha'].value_counts()
                
                summary_data = []
                for kode in sorted(kode_names.keys()):
                    count = value_counts.get(kode, 0)
                    if count > 0:
                        summary_data.append({
                            'Kode': kode,
                            'Bentuk Badan Hukum': kode_names[kode],
                            'Jumlah': count,
                            'Persentase': f"{count/len(df_proc)*100:.2f}%"
                        })
                
                unclassified = df_proc['bentuk_badan_hukum_usaha'].isna().sum()
                summary_data.append({
                    'Kode': '-',
                    'Bentuk Badan Hukum': 'Tidak Terklasifikasi',
                    'Jumlah': unclassified,
                    'Persentase': f"{unclassified/len(df_proc)*100:.2f}%"
                })
                
                df_summary = pd.DataFrame(summary_data)
                
                col1, col2 = st.columns([2, 1])
                
                with col1:
                    st.dataframe(df_summary, use_container_width=True, hide_index=True)
                
                with col2:
                    # Pie chart
                    import plotly.express as px
                    
                    chart_data = df_summary[df_summary['Jumlah'] > 0].copy()
                    fig = px.pie(
                        chart_data, 
                        values='Jumlah', 
                        names='Bentuk Badan Hukum',
                        title='Distribusi Bentuk Badan Hukum'
                    )
                    st.plotly_chart(fig, use_container_width=True)
                
                # Sample unclassified
                st.divider()
                st.subheader("📝 Sample Tidak Terklasifikasi")
                
                nama_col = 'nama_usaha' if 'nama_usaha' in df_proc.columns else df_proc.columns[1]
                unclassified_samples = df_proc[df_proc['bentuk_badan_hukum_usaha'].isna()][nama_col].drop_duplicates().head(20)
                
                if len(unclassified_samples) > 0:
                    st.write("Sample nama usaha yang belum terklasifikasi (untuk referensi menambah pattern):")
                    for i, nama in enumerate(unclassified_samples, 1):
                        st.text(f"{i:2d}. {nama}")
                else:
                    st.success("Semua data sudah terklasifikasi!")
                
                # Custom pattern
                st.divider()
                st.subheader("➕ Tambah Pattern Kustom")
                
                col1, col2, col3 = st.columns([1, 2, 1])
                
                with col1:
                    custom_kode = st.selectbox(
                        "Kode",
                        options=list(kode_names.keys()),
                        format_func=lambda x: f"{x} - {kode_names[x]}"
                    )
                
                with col2:
                    custom_pattern = st.text_input(
                        "Regex Pattern",
                        placeholder=r"(?i)\b(KEYWORD1|KEYWORD2)\b"
                    )
                
                with col3:
                    st.write("")
                    st.write("")
                    if st.button("Apply Pattern"):
                        if custom_pattern:
                            try:
                                mask_empty = df_proc['bentuk_badan_hukum_usaha'].isna()
                                
                                def apply_pattern(nama):
                                    if pd.isna(nama):
                                        return None
                                    if re.search(custom_pattern, str(nama).upper()):
                                        return custom_kode
                                    return None
                                
                                new_classified = df_proc.loc[mask_empty, nama_col].apply(apply_pattern)
                                count_new = new_classified.notna().sum()
                                
                                df_proc.loc[mask_empty, 'bentuk_badan_hukum_usaha'] = df_proc.loc[mask_empty, 'bentuk_badan_hukum_usaha'].fillna(new_classified)
                                st.session_state['df_processed'] = df_proc
                                
                                st.success(f"✅ {count_new:,} rows baru terklasifikasi sebagai kode {custom_kode}")
                                st.rerun()
                            except re.error as e:
                                st.error(f"Invalid regex pattern: {e}")
            else:
                st.info("Jalankan proses data di tab 'Processing' terlebih dahulu dengan opsi klasifikasi dicentang.")
        
        # Tab 4: Statistik
        with tab4:
            st.subheader("📈 Statistik Data")
            
            df_show = st.session_state.get('df_processed', df)
            
            col1, col2 = st.columns(2)
            
            with col1:
                st.write("**Info Kolom:**")
                col_info = []
                for col in df_show.columns:
                    col_info.append({
                        'Kolom': col,
                        'Tipe': str(df_show[col].dtype),
                        'Non-Null': f"{df_show[col].notna().sum():,}",
                        'Null': f"{df_show[col].isna().sum():,}",
                        'Unique': f"{df_show[col].nunique():,}"
                    })
                st.dataframe(pd.DataFrame(col_info), use_container_width=True, hide_index=True)
            
            with col2:
                st.write("**Memory Usage:**")
                mem_usage = df_show.memory_usage(deep=True).sum() / 1024**2
                st.metric("Total Memory", f"{mem_usage:.2f} MB")
                
                st.write("**Data Shape:**")
                st.metric("Rows", f"{len(df_show):,}")
                st.metric("Columns", len(df_show.columns))
                
                # Cek kolom alamat (bisa alamat, alamat_usaha, atau variasi lainnya)
                alamat_cols = [col for col in df_show.columns if 'alamat' in col.lower()]
                if alamat_cols:
                    alamat_col = alamat_cols[0]
                    st.write(f"**{alamat_col}:**")
                    filled = df_show[alamat_col].notna().sum()
                    st.metric("Filled", f"{filled:,} ({filled/len(df_show)*100:.1f}%)")
        
        # Tab 5: Download
        with tab5:
            st.subheader("💾 Download Hasil")
            
            df_download = st.session_state.get('df_processed', df)
            
            st.write(f"Data siap download: **{df_download.shape[0]:,} rows × {df_download.shape[1]} columns**")
            
            # Preview
            with st.expander("Preview Data"):
                st.dataframe(df_download.head(50), use_container_width=True)
            
            # Download buttons
            col1, col2 = st.columns(2)
            
            with col1:
                excel_data = to_excel_bytes(df_download)
                st.download_button(
                    label="📥 Download Excel (.xlsx)",
                    data=excel_data,
                    file_name=f"{output_filename}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            
            with col2:
                csv_data = to_csv_bytes(df_download)
                st.download_button(
                    label="📥 Download CSV (.csv)",
                    data=csv_data,
                    file_name=f"{output_filename}.csv",
                    mime="text/csv"
                )
    
    else:
        # No file uploaded
        st.info("👆 Upload file data usaha di sidebar untuk memulai")
        
        # Instructions
        with st.expander("📖 Panduan Penggunaan"):
            st.markdown("""
            ### Cara Menggunakan Aplikasi
            
            1. **Upload Data Usaha** (wajib)
               - Format: CSV atau Excel (.xlsx, .xls)
               - Harus memiliki kolom `nama_usaha` untuk klasifikasi
               
            2. **Upload Data Wilayah** (opsional)
               - Format: Excel (.xlsx, .xls)
               - Digunakan untuk mengisi kolom `alamat` yang kosong
               - Harus memiliki kolom: kdprov, kdkab, kdkec, kddesa, nmdesa, nmkec, nmkab, nmprov
               
            3. **Kode Kota**
               - Masukkan kode kabupaten yang sebenarnya adalah Kota
               - Default: 71 (Kota Gorontalo)
               - Pisahkan dengan koma untuk multiple kode
               
            4. **Proses Data**
               - Pilih langkah-langkah processing yang diinginkan
               - Klik tombol "Proses Data"
               
            5. **Download Hasil**
               - Pilih format output (Excel/CSV)
               - Klik tombol download
            
            ### Klasifikasi Bentuk Badan Hukum
            
            | Kode | Bentuk Badan Hukum | Pattern |
            |------|-------------------|---------|
            | 1 | Perseroan (PT) | PT, PERSEROAN |
            | 2 | Yayasan | YAYASAN |
            | 3 | Koperasi | KOPERASI, KOP, KUD |
            | 4 | Dana Pensiun | DANA PENSIUN |
            | 5 | Perum/Perumda | PERUM, PERUMDA |
            | 6 | BUM Desa | BUMDES |
            | 7 | CV | CV, C.V |
            | 8 | Firma | FIRMA, FA |
            | 9 | Persekutuan Perdata | PERSEKUTUAN |
            | 10 | Kantor Perwakilan LN | PERWAKILAN |
            | 11 | Badan Usaha LN | LLC, LTD, INC |
            | 12 | Usaha Perseorangan | UD, TOKO, WARUNG |
            """)


if __name__ == "__main__":
    main()
