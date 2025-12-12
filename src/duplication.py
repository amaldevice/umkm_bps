import pandas as pd
import numpy as np
import io
import re
from difflib import SequenceMatcher
from openpyxl import load_workbook

# ==========================================
# KONFIGURASI
# ==========================================
HEADER_ROW_PANDAS = 1
DATA_START_ROW_EXCEL = 3

COL_INDEX_IDSBR = 1
COL_INDEX_STATUS = 13
COL_INDEX_MASTER_ID = 14

THRESHOLD = 0.98

KAMUS_TYPO = {
    'JL': 'JALAN', 'JLN': 'JALAN', 'DSN': 'DUSUN', 'KP': 'KAMPUNG',
    'KEC': 'KECAMATAN', 'KAB': 'KABUPATEN', 'NO': 'NOMOR',
    'RT': '', 'RW': '', 'BLK': 'BLOK',
    'KOMPLEK': 'KOMPLEKS', 'WARUNG': 'TOKO', 'WR': 'TOKO',
    'UD': '', 'CV': '', 'PT': '', 'TB': 'TOKO'
}


def clean_text(text):
    if pd.isna(text): return ""
    text = str(text).upper()
    text = re.sub(r'[^\w\s]', ' ', text)
    words = text.split()
    fixed_words = [KAMUS_TYPO.get(w, w) for w in words]
    text = " ".join(fixed_words)
    text = re.sub(r'\s+', ' ', text).strip()
    return text


def get_consecutive_ranges(rows):
    """
    Mengubah list baris [2, 3, 4, 8, 9] menjadi ranges [(2, 3), (8, 2)]
    Format: (start_index, count)
    """
    if not rows:
        return []

    rows.sort()
    ranges = []
    start = rows[0]
    count = 1

    for i in range(1, len(rows)):
        if rows[i] == rows[i - 1] + 1:
            count += 1
        else:
            ranges.append((start, count))
            start = rows[i]
            count = 1
    ranges.append((start, count))
    return ranges


def process_duplication(uploaded_file):
    """
    Memproses file Excel untuk cek duplikasi.
    Optimasi: Menghapus baris unik menggunakan Batch Deletion agar tidak hang/stuck.
    """
    # Reset pointer
    uploaded_file.seek(0)

    # 1. BACA DATA
    df = pd.read_excel(uploaded_file, header=HEADER_ROW_PANDAS)

    # --- [LOGIKA SKOR MASTER] ---
    df_score = df.copy()
    df_score = df_score.replace(0, np.nan)
    df_score = df_score.replace(r'^[\s\-\.]*$', np.nan, regex=True)

    cols_to_score = [c for c in df_score.columns if c != 'idsbr']

    # Handle jika kolom kosong (jaga-jaga)
    if not cols_to_score:
        df['column_count_score'] = 0
        df['char_length_score'] = 0
    else:
        df['column_count_score'] = df_score[cols_to_score].notna().sum(axis=1)
        df['char_length_score'] = df_score[cols_to_score].fillna('').astype(str).apply(
            lambda x: x.str.len().sum(), axis=1)

    # --- 2. PEMBERSIHAN DATA ---
    if 'nama_usaha' not in df.columns or 'alamat' not in df.columns:
        raise ValueError("File harus memiliki kolom 'nama_usaha' dan 'alamat'")

    df['nama_clean'] = df['nama_usaha'].apply(clean_text)
    df['alamat_clean'] = df['alamat'].apply(clean_text)
    df['sort_key'] = df['nama_clean'] + " " + df['alamat_clean']

    # --- 3. LOGIKA FUZZY 98% ---
    df = df.sort_values(by=['sort_key', 'idsbr'])

    df['group_id'] = -1
    records = df.to_dict('records')

    current_group = 0
    prev_nama = ""
    prev_alamat = ""
    grouped_records = []

    # Logic Loop Fuzzy
    for i, row in enumerate(records):
        curr_nama = row['nama_clean']
        curr_alamat = row['alamat_clean']

        if i == 0:
            row['group_id'] = current_group
            prev_nama = curr_nama
            prev_alamat = curr_alamat
        else:
            if prev_nama == "" or curr_nama == "":
                ratio_nama = 0
            elif prev_nama == curr_nama:
                ratio_nama = 1.0
            else:
                ratio_nama = SequenceMatcher(None, prev_nama, curr_nama).ratio()

            if prev_alamat == "" or curr_alamat == "":
                ratio_alamat = 0
            elif prev_alamat == curr_alamat:
                ratio_alamat = 1.0
            else:
                ratio_alamat = SequenceMatcher(None, prev_alamat, curr_alamat).ratio()

            if ratio_nama >= THRESHOLD and ratio_alamat >= THRESHOLD:
                row['group_id'] = current_group
            else:
                current_group += 1
                row['group_id'] = current_group
                prev_nama = curr_nama
                prev_alamat = curr_alamat

        grouped_records.append(row)

    df_processed = pd.DataFrame(grouped_records)

    # --- 4. PENENTUAN MASTER ---
    df_processed = df_processed.sort_values(
        by=['group_id', 'column_count_score', 'char_length_score', 'idsbr'],
        ascending=[True, False, False, False]
    )

    df_processed['master_id_final'] = df_processed.groupby('group_id')['idsbr'].transform('first')

    def get_final_status(row):
        if not row['nama_clean']: return None, None
        if row['idsbr'] == row['master_id_final']:
            return 1, None
        else:
            return 9, row['master_id_final']

    results = df_processed.apply(get_final_status, axis=1, result_type='expand')
    df_processed['final_status'] = results[0]
    df_processed['final_master_id'] = results[1]

    # --- 5. FILTER DUPLIKASI ---
    group_counts = df_processed['group_id'].value_counts()
    duplicate_group_ids = group_counts[group_counts > 1].index

    # Set ID yang harus disimpan (Duplikat & Masternya)
    ids_to_keep = set(df_processed[df_processed['group_id'].isin(duplicate_group_ids)]['idsbr'].astype(str))

    data_map = {}
    for _, row in df_processed.iterrows():
        str_id = str(row['idsbr']).strip()
        if str_id in ids_to_keep:
            data_map[str_id] = {'status': row['final_status'], 'master': row['final_master_id']}

    # --- 6. TULIS KE EXCEL (OPTIMIZED BATCH DELETION) ---
    uploaded_file.seek(0)
    book = load_workbook(uploaded_file)
    sheet = book.active

    rows_to_delete = []

    # Iterasi untuk Update Data sekaligus menandai baris yang dihapus
    for row in sheet.iter_rows(min_row=DATA_START_ROW_EXCEL, max_col=COL_INDEX_MASTER_ID):
        # Ambil value ID
        cell_id = row[COL_INDEX_IDSBR - 1]
        val_id = str(cell_id.value).strip() if cell_id.value is not None else ""

        if val_id in ids_to_keep:
            # Jika ini bagian duplikasi, UPDATE kolom status & master
            if val_id in data_map:
                vals = data_map[val_id]
                if vals['status'] is not None:
                    row[COL_INDEX_STATUS - 1].value = vals['status']
                    val_master = vals['master']
                    if pd.isna(val_master): val_master = None
                    row[COL_INDEX_MASTER_ID - 1].value = val_master
        else:
            # Jika ini unik (tidak duplikasi), TANDAI untuk dihapus
            # row[0].row adalah nomor baris Excel (1-based)
            rows_to_delete.append(row[0].row)

    # --- BAGIAN PENTING: MENGHAPUS MENGGUNAKAN RANGE ---
    # Jika kita hapus satu per satu, Excel akan 'hang' karena shifting ribuan kali.
    # Kita ubah daftar baris [5, 6, 7, 10] menjadi range [(5,3), (10,1)]
    # Lalu hapus dari bawah ke atas.

    if rows_to_delete:
        deletion_ranges = get_consecutive_ranges(rows_to_delete)

        # Hapus secara terbalik agar index tidak bergeser untuk range di atasnya
        for start_idx, count in reversed(deletion_ranges):
            sheet.delete_rows(start_idx, amount=count)

    output = io.BytesIO()
    book.save(output)
    output.seek(0)

    return output