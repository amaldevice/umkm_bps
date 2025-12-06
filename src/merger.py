import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows


def process_merge_files(uploaded_files, header_row=1):
    """
    Menggabungkan multiple file Excel dengan mempertahankan format file pertama.

    Args:
        uploaded_files: List file dari Streamlit
        header_row: Index baris header (0-based).
                    Jika di Excel baris 2, maka inputnya 1.
    """
    if not uploaded_files:
        raise ValueError("Tidak ada file yang diupload.")

    all_dfs = []

    # 1. BACA SEMUA DATA
    for file in uploaded_files:
        file.seek(0)  # Reset pointer file
        try:
            # Baca data
            df = pd.read_excel(file, header=header_row)
            all_dfs.append(df)
        except Exception as e:
            print(f"Skip file {file.name}: {e}")
            continue

    if not all_dfs:
        raise ValueError("Gagal membaca data dari file yang diupload.")

    # 2. GABUNGKAN DATA
    merged_df = pd.concat(all_dfs, ignore_index=True)
    total_rows = len(merged_df)

    # 3. SIAPKAN TEMPLATE (Gunakan file pertama sebagai base)
    # Kita akan menulis ulang data ke dalam file pertama agar format header terjaga
    base_file = uploaded_files[0]
    base_file.seek(0)

    # Load workbook menggunakan openpyxl
    wb = load_workbook(base_file)
    ws = wb.active  # Ambil sheet aktif

    # Tentukan baris mulai menulis data
    # header_row adalah index (misal 1).
    # Di Excel itu Row 2. Berarti data mulai di Row 3.
    start_row = header_row + 2

    # 4. TULIS DATA KE TEMPLATE
    # Kita gunakan iterasi rows untuk menulis value saja, agar style cell tidak rusak total
    # (Meskipun baris baru di bawah mungkin tidak punya border, header tetap aman)

    rows = dataframe_to_rows(merged_df, index=False, header=False)

    for r_idx, row in enumerate(rows, 1):
        for c_idx, value in enumerate(row, 1):
            # Tulis value ke posisi yang tepat
            # row = start_row + urutan data
            ws.cell(row=start_row + r_idx - 1, column=c_idx, value=value)

    # 5. SIMPAN HASIL
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output, total_rows