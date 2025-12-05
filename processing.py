import pandas as pd
import re
from typing import Optional, List, Dict
from io import BytesIO
import warnings
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

warnings.filterwarnings('ignore')

# =============================================================================
# HELPER FUNCTIONS
# =============================================================================

def load_data(uploaded_file, skiprows=0):
    """Load data dari uploaded file (CSV atau Excel)."""
    # If uploaded_file is a file path or file-like object
    if hasattr(uploaded_file, 'filename'):
        file_name = uploaded_file.filename.lower()
    elif hasattr(uploaded_file, 'name'):
        file_name = uploaded_file.name.lower()
    else:
        file_name = "unknown"
        
    # Pastikan skiprows >= 0
    skiprows = max(0, int(skiprows)) if skiprows else 0
    
    if file_name.endswith(('.xlsx', '.xls')):
        df = pd.read_excel(uploaded_file, skiprows=skiprows)
    else:
        # Coba berbagai encoding untuk CSV
        try:
            df = pd.read_csv(uploaded_file, encoding='utf-8', skiprows=skiprows)
        except UnicodeDecodeError:
            try:
                uploaded_file.seek(0)
                df = pd.read_csv(uploaded_file, encoding='latin-1', skiprows=skiprows)
            except UnicodeDecodeError:
                uploaded_file.seek(0)
                df = pd.read_csv(uploaded_file, encoding='cp1252', skiprows=skiprows)
    
    return df


def clean_columns(df: pd.DataFrame) -> tuple:
    """Bersihkan nama kolom."""
    df = df.copy()
    original_columns = df.columns.tolist()
    cleaned_columns = df.columns.str.split('\n').str[0].tolist()
    
    # Buat mapping dari cleaned -> original
    original_columns_mapping = dict(zip(cleaned_columns, original_columns))
    
    df.columns = cleaned_columns
    return df, original_columns_mapping


def restore_original_columns(df: pd.DataFrame, original_columns_mapping: dict) -> pd.DataFrame:
    """Kembalikan nama kolom ke format original."""
    df = df.copy()
    # Rename kolom yang ada di mapping
    rename_map = {col: original_columns_mapping.get(col, col) for col in df.columns}
    df = df.rename(columns=rename_map)
    return df


def format_kode_wilayah(df: pd.DataFrame) -> pd.DataFrame:
    """Format kode wilayah dengan zero-padding."""
    df = df.copy()
    
    if 'kdprov' in df.columns:
        df['kdprov'] = df['kdprov'].astype("Int64").astype(str).str.zfill(2)
    if 'kdkab' in df.columns:
        df['kdkab'] = df['kdkab'].astype("Int64").astype(str).str.zfill(2)
    if 'kdkec' in df.columns:
        df['kdkec'] = df['kdkec'].astype("Int64").astype(str).str.zfill(3)
    if 'kddesa' in df.columns:
        df['kddesa'] = df['kddesa'].astype("Int64").astype(str).str.zfill(3)
    
    return df


def format_wilayah_lookup(df: pd.DataFrame) -> pd.DataFrame:
    """Format kode wilayah di lookup table."""
    df = df.copy()
    
    df['kdprov'] = df['kdprov'].astype("Int64").astype(str).str.zfill(2)    
    df['kdkab'] = df['kdkab'].astype("Int64").astype(str).str.zfill(2)
    df['kdkec'] = df['kdkec'].astype("Int64").astype(str).str.zfill(3)
    df['kddesa'] = df['kddesa'].astype("Int64").astype(str).str.zfill(3)
    
    return df


def create_alamat_new(df_wilayah: pd.DataFrame, kota_codes: List[str]) -> pd.DataFrame:
    """Buat kolom alamat_new dari data wilayah (untuk lookup table)."""
    df = df_wilayah.copy()
    
    # Pastikan kolom nama bertipe string
    for col in ["nmdesa", "nmkec", "nmkab", "nmprov"]:
        if col in df.columns:
            df[col] = df[col].fillna('').astype(str).str.strip()
    
    def build_alamat(row):
        parts = []
        
        if row.get('nmdesa', '') and row['nmdesa'] not in ['', 'nan', 'None', '<NA>']:
            parts.append(f"Desa {row['nmdesa']}")
        
        if row.get('nmkec', '') and row['nmkec'] not in ['', 'nan', 'None', '<NA>']:
            parts.append(f"Kecamatan {row['nmkec']}")
        
        if row.get('nmkab', '') and row['nmkab'] not in ['', 'nan', 'None', '<NA>']:
            kdkab = str(row.get('kdkab', ''))
            prefix = "Kota" if kdkab in kota_codes else "Kabupaten"
            parts.append(f"{prefix} {row['nmkab']}")
        
        if row.get('nmprov', '') and row['nmprov'] not in ['', 'nan', 'None', '<NA>']:
            parts.append(f"Provinsi {row['nmprov']}")
        
        return ', '.join(parts) if parts else None
    
    df['alamat_new'] = df.apply(build_alamat, axis=1)
    return df


def create_fallback_alamat(df: pd.DataFrame, df_wilayah: pd.DataFrame, kota_codes: List[str], alamat_col: str = 'alamat') -> pd.DataFrame:
    """Buat alamat fallback untuk rows yang tidak match. Isi kolom alamat yang sudah ada jika kosong."""
    df = df.copy()
    
    # Pastikan kolom alamat ada
    if alamat_col not in df.columns:
        df[alamat_col] = None
    
    # Hanya isi yang kosong/null/NaN
    mask_empty = df[alamat_col].isna()
    if mask_empty.any():
        # Cek juga yang string kosong atau 'nan'
        mask_str_empty = ~df[alamat_col].isna() & (df[alamat_col].astype(str).str.strip().isin(['', 'nan', 'None', '<NA>']))
        mask_empty = mask_empty | mask_str_empty
    
    if not mask_empty.any():
        return df
    
    # Lookup tables
    lookup_kec = df_wilayah[['kdprov', 'kdkab', 'kdkec', 'nmkec', 'nmkab', 'nmprov']].drop_duplicates(
        subset=['kdprov', 'kdkab', 'kdkec']
    )
    lookup_kab = df_wilayah[['kdprov', 'kdkab', 'nmkab', 'nmprov']].drop_duplicates(
        subset=['kdprov', 'kdkab']
    )
    lookup_prov = df_wilayah[['kdprov', 'nmprov']].drop_duplicates(subset=['kdprov'])
    
    def build_fallback(row):
        parts = []
        kdprov = str(row.get('kdprov', ''))
        kdkab = str(row.get('kdkab', ''))
        kdkec = str(row.get('kdkec', ''))
        
        # Coba match level kecamatan
        match_kec = lookup_kec[
            (lookup_kec['kdprov'] == kdprov) & 
            (lookup_kec['kdkab'] == kdkab) & 
            (lookup_kec['kdkec'] == kdkec)
        ]
        
        if not match_kec.empty:
            r = match_kec.iloc[0]
            if r.get('nmkec', '') not in ['', 'nan', 'None']:
                parts.append(f"Kecamatan {r['nmkec']}")
            if r.get('nmkab', '') not in ['', 'nan', 'None']:
                prefix = "Kota" if kdkab in kota_codes else "Kabupaten"
                parts.append(f"{prefix} {r['nmkab']}")
            if r.get('nmprov', '') not in ['', 'nan', 'None']:
                parts.append(f"Provinsi {r['nmprov']}")
            if parts:
                return ', '.join(parts)
        
        # Coba match level kabupaten
        match_kab = lookup_kab[
            (lookup_kab['kdprov'] == kdprov) & 
            (lookup_kab['kdkab'] == kdkab)
        ]
        
        if not match_kab.empty:
            r = match_kab.iloc[0]
            if r.get('nmkab', '') not in ['', 'nan', 'None']:
                prefix = "Kota" if kdkab in kota_codes else "Kabupaten"
                parts.append(f"{prefix} {r['nmkab']}")
            if r.get('nmprov', '') not in ['', 'nan', 'None']:
                parts.append(f"Provinsi {r['nmprov']}")
            if parts:
                return ', '.join(parts)
        
        # Coba match level provinsi
        match_prov = lookup_prov[lookup_prov['kdprov'] == kdprov]
        if not match_prov.empty:
            nmprov = match_prov.iloc[0].get('nmprov', '')
            if nmprov not in ['', 'nan', 'None']:
                return f"Provinsi {nmprov}"
        
        # Fallback ke kode
        codes = []
        if kdprov: codes.append(f"Prov:{kdprov}")
        if kdkab: codes.append(f"Kab:{kdkab}")
        if kdkec: codes.append(f"Kec:{kdkec}")
        kddesa = str(row.get('kddesa', ''))
        if kddesa: codes.append(f"Desa:{kddesa}")
        
        return f"[Kode: {', '.join(codes)}]" if codes else None
    
    # Isi kolom alamat yang kosong dengan fallback
    df.loc[mask_empty, alamat_col] = df.loc[mask_empty].apply(build_fallback, axis=1)
    return df


def merge_with_wilayah(df: pd.DataFrame, df_wilayah: pd.DataFrame, kota_codes: List[str], alamat_col: str = 'alamat') -> pd.DataFrame:
    """Merge data usaha dengan data wilayah. Isi kolom alamat yang sudah ada jika kosong."""
    df = df.copy()
    
    # Pastikan kolom alamat ada
    if alamat_col not in df.columns:
        df[alamat_col] = None
    
    df_wilayah = format_wilayah_lookup(df_wilayah)
    df_wilayah = create_alamat_new(df_wilayah, kota_codes)
    
    on_columns = ["kdprov", "kdkab", "kdkec", "kddesa"]
    lookup = df_wilayah[on_columns + ["alamat_new"]].drop_duplicates(subset=on_columns)
    
    # Merge untuk mendapatkan alamat_new dari lookup
    df = df.merge(lookup, on=on_columns, how='left')
    
    # Isi kolom alamat yang kosong dengan alamat_new
    mask_empty = df[alamat_col].isna()
    if mask_empty.any():
        # Cek juga yang string kosong atau 'nan'
        mask_str_empty = ~df[alamat_col].isna() & (df[alamat_col].astype(str).str.strip().isin(['', 'nan', 'None', '<NA>']))
        mask_empty = mask_empty | mask_str_empty
    
    df.loc[mask_empty & df['alamat_new'].notna(), alamat_col] = df.loc[mask_empty & df['alamat_new'].notna(), 'alamat_new']
    
    # Hapus kolom alamat_new yang temporary
    if 'alamat_new' in df.columns:
        df = df.drop(columns=['alamat_new'])
    
    # Fallback untuk yang masih kosong
    df = create_fallback_alamat(df, df_wilayah, kota_codes, alamat_col)
    
    return df


def get_bentuk_badan_hukum_patterns() -> dict:
    """Definisi pattern regex untuk klasifikasi bentuk badan hukum."""
    patterns = {
        1: r'(?i)\b(PT\.?|P\.T\.?|PERSEROAN|PERSERO)\b',
        2: r'(?i)\b(YAYASAN|YYS\.?)\b',
        3: r'(?i)\b(KOPERASI|KOP\.?|KOPKAR|KSU|KSP|KOPWAN|KOPDIT|KOSPIN|KUD)\b',
        4: r'(?i)\b(DANA\s*PENSIUN|DAPEN|DP\.?)\b',
        5: r'(?i)\b(PERUM|PERUMDA|PERUMNAS)\b',
        6: r'(?i)\b(BUMDES|BUM\s*DESA|BUMN?\s*DES)\b',
        7: r'(?i)\b(CV\.?|C\.V\.?)\b',
        8: r'(?i)\b(FIRMA|FA\.?)\b',
        9: r'(?i)\b(PERSEKUTUAN|MAATSCHAP|SEKUTU)\b',
        10: r'(?i)\b(PERWAKILAN|REPRESENTATIVE\s*OFFICE|REP\.?\s*OFFICE)\b',
        11: r'(?i)\b(LLC|LTD\.?|LIMITED|INC\.?|INCORPORATED|CORP\.?|CORPORATION|PTE\.?\s*LTD\.?|SDN\.?\s*BHD\.?|GMBH|AG\b|S\.?A\.?\b|BV\.?|NV\.?)\b',
        12: r'(?i)\b(UD\.?|U\.D\.?|PD\.?|P\.D\.?|TOKO|WARUNG|KIOS|KEDAI|DEPOT|SALON|BENGKEL|LAUNDRY|FOTOCOPY|FOTO\s*COPY|PERCETAKAN|SABLON|KONTER|COUNTER|RENTAL|SEWA|JAHIT|TAILOR|CATERING|KATERING)\b',
    }
    return patterns


def classify_bentuk_badan_hukum(df: pd.DataFrame, nama_col: str = 'nama_usaha') -> pd.DataFrame:
    """Klasifikasi bentuk badan hukum berdasarkan nama usaha. Overwrite kolom bentuk_badan_hukum_usaha."""
    df = df.copy()
    patterns = get_bentuk_badan_hukum_patterns()
    
    # Pastikan kolom bentuk_badan_hukum_usaha ada
    if 'bentuk_badan_hukum_usaha' not in df.columns:
        df['bentuk_badan_hukum_usaha'] = None
    
    def classify_single(nama):
        if pd.isna(nama) or str(nama).strip() == '':
            return None
        nama_str = str(nama).upper()
        for kode, pattern in patterns.items():
            if re.search(pattern, nama_str):
                return kode
        return None
    
    # Overwrite kolom bentuk_badan_hukum_usaha
    df['bentuk_badan_hukum_usaha'] = df[nama_col].apply(classify_single)
    return df


def get_kode_names() -> dict:
    """Mapping kode ke nama bentuk badan hukum."""
    return {
        1: 'Perseroan (PT)',
        2: 'Yayasan',
        3: 'Koperasi',
        4: 'Dana Pensiun',
        5: 'Perum/Perumda',
        6: 'BUM Desa',
        7: 'CV',
        8: 'Firma',
        9: 'Persekutuan Perdata',
        10: 'Kantor Perwakilan LN',
        11: 'Badan Usaha LN',
        12: 'Usaha Perseorangan',
        99: 'Lainnya'
    }


def detect_and_fill_jaringan_usaha(df: pd.DataFrame, nama_col: str = 'nama_usaha') -> pd.DataFrame:
    """Deteksi format <XXXX> dalam nama usaha dan isi bentuk_badan_hukum_usaha=12 dan jaringan_usaha=1."""
    df = df.copy()
    
    # Pastikan kolom ada
    if 'bentuk_badan_hukum_usaha' not in df.columns:
        df['bentuk_badan_hukum_usaha'] = None
    
    if 'jaringan_usaha' not in df.columns:
        df['jaringan_usaha'] = None
    
    # Pastikan kolom nama_usaha ada
    if nama_col not in df.columns:
        return df
    
    # Pattern untuk mendeteksi format <XXXX>
    pattern = r'<[^>]+>'
    
    # Deteksi format <XXXX> pada semua nama usaha
    def has_jaringan_format(nama):
        if pd.isna(nama):
            return False
        
        nama_str = str(nama).strip()
        
        # Skip jika kosong atau invalid
        if nama_str in ['', 'nan', 'None', '<NA>']:
            return False
        
        # Deteksi format <XXXX>
        return bool(re.search(pattern, nama_str))
    
    # Buat mask untuk rows yang memiliki format <XXXX>
    mask_jaringan = df[nama_col].apply(has_jaringan_format)
    
    # Set bentuk_badan_hukum_usaha = 12 untuk rows yang terdeteksi
    df.loc[mask_jaringan, 'bentuk_badan_hukum_usaha'] = 12
    
    # Set jaringan_usaha = 1 untuk rows yang terdeteksi
    df.loc[mask_jaringan, 'jaringan_usaha'] = 1
    
    return df


def fill_nomor_whatsapp(df: pd.DataFrame, nomor_telepon_col: str = 'nomor_telepon', nomor_whatsapp_col: str = 'nomor_whatsapp') -> pd.DataFrame:
    """Isi kolom nomor_whatsapp jika kosong dari nomor_telepon dengan format 08... diubah menjadi +628..."""
    df = df.copy()
    
    # Pastikan kolom nomor_whatsapp ada
    if nomor_whatsapp_col not in df.columns:
        df[nomor_whatsapp_col] = None
    
    # Pastikan kolom nomor_telepon ada
    if nomor_telepon_col not in df.columns:
        return df
    
    # Fungsi untuk convert nomor telepon ke format whatsapp
    def convert_to_whatsapp(nomor_telepon):
        if pd.isna(nomor_telepon):
            return None
        
        # Convert ke string dan strip whitespace
        nomor_str = str(nomor_telepon).strip()
        
        # Cek jika kosong atau invalid
        if nomor_str in ['', 'nan', 'None', '<NA>']:
            return None
        
        # Cek jika format dimulai dengan 08
        if nomor_str.startswith('08'):
            # Ubah format dari 08... menjadi +628...
            # Hapus karakter non-digit terlebih dahulu untuk memastikan hanya angka
            nomor_clean = re.sub(r'\D', '', nomor_str)
            if nomor_clean.startswith('08'):
                return '+62' + nomor_clean[1:]  # Ganti 0 dengan +62
        
        return None
    
    # Cari rows yang nomor_whatsapp kosong
    mask_empty_whatsapp = df[nomor_whatsapp_col].isna()
    if mask_empty_whatsapp.any():
        # Cek juga yang string kosong atau 'nan'
        mask_str_empty = ~df[nomor_whatsapp_col].isna() & (
            df[nomor_whatsapp_col].astype(str).str.strip().isin(['', 'nan', 'None', '<NA>'])
        )
        mask_empty_whatsapp = mask_empty_whatsapp | mask_str_empty
    
    if not mask_empty_whatsapp.any():
        return df
    
    # Apply conversion untuk rows yang nomor_whatsapp kosong dan nomor_telepon ada
    mask_has_telepon = df[nomor_telepon_col].notna()
    mask_to_fill = mask_empty_whatsapp & mask_has_telepon
    
    if mask_to_fill.any():
        # Convert nomor_telepon ke format whatsapp
        converted = df.loc[mask_to_fill, nomor_telepon_col].apply(convert_to_whatsapp)
        
        # Hanya isi yang berhasil dikonversi (tidak None)
        mask_valid = converted.notna()
        if mask_valid.any():
            df.loc[mask_to_fill & mask_valid, nomor_whatsapp_col] = converted[mask_valid]
    
    return df


def clean_nama_usaha(df: pd.DataFrame, nama_col: str = 'nama_usaha') -> pd.DataFrame:
    """Clean penamaan usaha di kolom nama_usaha dengan regex."""
    df = df.copy()
    
    # Pastikan kolom nama_usaha ada
    if nama_col not in df.columns:
        return df
    
    def clean_single_nama(nama):
        if pd.isna(nama):
            return None
        
        nama_str = str(nama).strip()
        
        # Skip jika kosong atau invalid
        if nama_str in ['', 'nan', 'None', '<NA>']:
            return nama_str
        
        # Cek apakah sudah dalam format yang benar
        sudah_benar = re.search(r',\s*(PT|CV|UD|PD|FA|P\.T\.?|C\.V\.?|U\.D\.?|P\.D\.?|F\.A\.?)\s*$', nama_str, re.IGNORECASE)
        skip_badan_hukum = sudah_benar is not None
        
        # Step 1: Ubah semua jenis kurung menjadi <>
        nama_str = re.sub(r'\(([^)]*)\)', r'<\1>', nama_str)
        nama_str = re.sub(r'\[([^\]]*)\]', r'<\1>', nama_str)
        nama_str = re.sub(r'\{([^}]*)\}', r'<\1>', nama_str)
        nama_str = re.sub(r'["""]([^"""]*)["""]', r'<\1>', nama_str)
        nama_str = re.sub(r"['']([^'']*)['']", r'<\1>', nama_str)
        nama_str = re.sub(r'`([^`]*)`', r'<\1>', nama_str)
        
        if skip_badan_hukum:
            nama_str = re.sub(r'\s+', ' ', nama_str).strip()
            nama_str = re.sub(r',\s*,\s*(PT|CV|UD|PD|FA|P\.T\.?|C\.V\.?|U\.D\.?|P\.D\.?|F\.A\.?)\s*$', r', \1', nama_str, flags=re.IGNORECASE)
            nama_str = re.sub(r',{2,}', ',', nama_str)
            nama_str = re.sub(r',\s*,', ', ', nama_str)
            return nama_str
        
        # Step 2: Memindahkan bentuk badan hukum dari depan ke belakang
        badan_hukum_patterns = [
            (r'^PT\.?\s+', 'PT'),
            (r'^P\.T\.?\s+', 'PT'),
            (r'^CV\.?\s+', 'CV'),
            (r'^C\.V\.?\s+', 'CV'),
            (r'^UD\.?\s+', 'UD'),
            (r'^U\.D\.?\s+', 'UD'),
            (r'(?<!S\.)\bPD\.?\s+', 'PD'),
            (r'(?<!S\.)\bP\.D\.?\s+', 'PD'),
            (r'^FA\.?\s+', 'FA'),
            (r'^F\.A\.?\s+', 'FA'),
        ]
        
        moved_from_front = False
        for pattern, suffix in badan_hukum_patterns:
            match = re.match(pattern, nama_str, re.IGNORECASE)
            if match:
                nama_str = re.sub(pattern, '', nama_str, flags=re.IGNORECASE).strip()
                if not re.search(r',\s*(PT|CV|UD|PD|FA|P\.T\.?|C\.V\.?|U\.D\.?|P\.D\.?|F\.A\.?)\s*$', nama_str, re.IGNORECASE):
                    nama_str = f"{nama_str}, {suffix}"
                moved_from_front = True
                break
        
        # Step 3: Menangani bentuk badan hukum di belakang dengan titik
        if not moved_from_front:
            if not re.search(r',\s*(PT|CV|UD|PD|FA|P\.T\.?|C\.V\.?|U\.D\.?|P\.D\.?|F\.A\.?)\s*$', nama_str, re.IGNORECASE):
                nama_str = re.sub(r'\.\s*PT\.?\s*$', ', PT', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\.\s*P\.T\.?\s*$', ', PT', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+PT\.?\s*$', ', PT', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+P\.T\.?\s*$', ', PT', nama_str, flags=re.IGNORECASE)
                
                nama_str = re.sub(r'\.\s*CV\.?\s*$', ', CV', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\.\s*C\.V\.?\s*$', ', CV', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+CV\.?\s*$', ', CV', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+C\.V\.?\s*$', ', CV', nama_str, flags=re.IGNORECASE)
                
                nama_str = re.sub(r'\.\s*UD\.?\s*$', ', UD', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\.\s*U\.D\.?\s*$', ', UD', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+UD\.?\s*$', ', UD', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+U\.D\.?\s*$', ', UD', nama_str, flags=re.IGNORECASE)
                
                nama_str = re.sub(r'\.\s*PD\.?\s*$', ', PD', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\.\s*P\.D\.?\s*$', ', PD', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+PD\.?\s*$', ', PD', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+P\.D\.?\s*$', ', PD', nama_str, flags=re.IGNORECASE)
                
                nama_str = re.sub(r'\.\s*FA\.?\s*$', ', FA', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\.\s*F\.A\.?\s*$', ', FA', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+FA\.?\s*$', ', FA', nama_str, flags=re.IGNORECASE)
                nama_str = re.sub(r'\s+F\.A\.?\s*$', ', FA', nama_str, flags=re.IGNORECASE)
        
        nama_str = re.sub(r'\s+', ' ', nama_str).strip()
        nama_str = re.sub(r',\s*,\s*(PT|CV|UD|PD|FA|P\.T\.?|C\.V\.?|U\.D\.?|P\.D\.?|F\.A\.?)\s*$', r', \1', nama_str, flags=re.IGNORECASE)
        nama_str = re.sub(r',{2,}', ',', nama_str)
        nama_str = re.sub(r',\s*,', ', ', nama_str)
        
        return nama_str
    
    df[nama_col] = df[nama_col].apply(clean_single_nama)
    return df


def organize_columns_by_category(df: pd.DataFrame) -> tuple:
    """Organisir kolom berdasarkan kategori yang ditentukan."""
    df = df.copy()
    
    categories = {
        'IDENTITAS USAHA/PERUSAHAAN': [
            'idsbr', 'nama_usaha', 'nama_komersial_usaha', 'alamat', 'nama_sls',
            'kodepos', 'nomor_te', 'nomor_wemail', 'website', 'latitude', 'longitude'
        ],
        'KEBERADAAN USAHA/PERUSAHAAN': [
            'keberadaan_usaha', 'keberadaan_usaha/perusahaan', 'idsbr_master',
            'kdprov_pindah', 'kdkab_pindah', 'kdkec_pindah', 'kddesa_pindah'
        ],
        'WILAYAH USAHA/PERUSAHAAN': [
            'kdprov', 'kdkab', 'kdkec', 'kddesa', 'nmprov', 'nmkab', 'nmkec', 'nmdesa'
        ],
        'KARAKTERISTIK USAHA/PERUSAHAAN': [
            'bentuk_badan_hukum_usaha', 'status_usaha', 'jenis_usaha',
            'kategori_usaha', 'skala_usaha', 'modal_usaha'
        ],
        'KEGIATAN USAHA/PERUSAHAAN': [
            'kode_kbli', 'nama_kbli', 'kegiatan_utama', 'kegiatan_tambahan'
        ],
        'LAIN LAIN': []
    }
    
    ordered_columns = []
    category_ranges = []
    
    def get_category_columns(category_name, max_cols):
        cat_cols = []
        for col in categories[category_name]:
            if col in df.columns and col not in ordered_columns and len(cat_cols) < max_cols:
                cat_cols.append(col)
        if len(cat_cols) < max_cols:
            for col in df.columns:
                if col in ordered_columns:
                    continue
                col_lower = col.lower()
                for pattern in categories[category_name]:
                    pattern_lower = pattern.lower()
                    if (pattern_lower in col_lower or col_lower in pattern_lower) and len(cat_cols) < max_cols:
                        cat_cols.append(col)
                        break
        return cat_cols
    
    identitas_cols = get_category_columns('IDENTITAS USAHA/PERUSAHAAN', 12)
    start_idx = len(ordered_columns)
    ordered_columns.extend(identitas_cols)
    if identitas_cols:
        category_ranges.append(('IDENTITAS USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    keberadaan_cols = get_category_columns('KEBERADAAN USAHA/PERUSAHAAN', 4)
    start_idx = len(ordered_columns)
    ordered_columns.extend(keberadaan_cols)
    if keberadaan_cols:
        category_ranges.append(('KEBERADAAN USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    wilayah_cols = get_category_columns('WILAYAH USAHA/PERUSAHAAN', 4)
    start_idx = len(ordered_columns)
    ordered_columns.extend(wilayah_cols)
    if wilayah_cols:
        category_ranges.append(('WILAYAH USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    karakteristik_cols = get_category_columns('KARAKTERISTIK USAHA/PERUSAHAAN', 6)
    start_idx = len(ordered_columns)
    ordered_columns.extend(karakteristik_cols)
    if karakteristik_cols:
        category_ranges.append(('KARAKTERISTIK USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    kegiatan_cols = get_category_columns('KEGIATAN USAHA/PERUSAHAAN', 4)
    start_idx = len(ordered_columns)
    ordered_columns.extend(kegiatan_cols)
    if kegiatan_cols:
        category_ranges.append(('KEGIATAN USAHA/PERUSAHAAN', start_idx, len(ordered_columns) - 1))
    
    remaining_cols = [col for col in df.columns if col not in ordered_columns][:2]
    start_idx = len(ordered_columns)
    ordered_columns.extend(remaining_cols)
    if remaining_cols:
        category_ranges.append(('LAIN LAIN', start_idx, len(ordered_columns) - 1))
    
    all_remaining = [col for col in df.columns if col not in ordered_columns]
    ordered_columns.extend(all_remaining)
    
    df_ordered = df[ordered_columns]
    
    return df_ordered, ordered_columns, category_ranges


def to_excel_bytes(df: pd.DataFrame, original_columns_mapping: Optional[Dict[str, str]] = None) -> bytes:
    """Convert DataFrame ke Excel bytes dengan format terorganisir."""
    output = BytesIO()
    
    if original_columns_mapping:
        df = restore_original_columns(df, original_columns_mapping)
    
    df_ordered = df.copy()
    ordered_columns = list(df_ordered.columns)
    total_cols = len(ordered_columns)
    
    segments = [
        ('IDENTITAS USAHA/PERUSAHAAN', 12),
        ('KEBERADAAN USAHA/PERUSAHAAN', 4),
        ('WILAYAH USAHA/PERUSAHAAN', 4),
        ('KARAKTERISTIK USAHA/PERUSAHAAN', 6),
        ('KEGIATAN USAHA/PERUSAHAAN', 4),
        ('LAIN LAIN', 2),
    ]
    category_ranges = []
    current_idx = 0
    for cat_name, width in segments:
        if current_idx >= total_cols:
            break
        start_idx = current_idx
        end_idx = min(current_idx + width - 1, total_cols - 1)
        category_ranges.append((cat_name, start_idx, end_idx))
        current_idx = end_idx + 1
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_ordered.to_excel(writer, index=False, sheet_name='Sheet1', startrow=1)
        
        workbook = writer.book
        worksheet = writer.sheets['Sheet1']
        
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        border_style = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for cat_name, start_idx, end_idx in category_ranges:
            start_letter = get_column_letter(start_idx + 1)
            end_letter = get_column_letter(end_idx + 1)
            
            if start_idx < end_idx:
                worksheet.merge_cells(f'{start_letter}1:{end_letter}1')
            
            cell = worksheet[f'{start_letter}1']
            cell.value = cat_name
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border_style
            
            if start_idx == end_idx:
                cell.border = border_style
        
        header_row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        header_row_font = Font(bold=True, size=10)
        
        for col_idx, col_name in enumerate(ordered_columns, 1):
            col_letter = get_column_letter(col_idx)
            cell = worksheet[f'{col_letter}2']
            cell.fill = header_row_fill
            cell.font = header_row_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_style
        
        worksheet.row_dimensions[1].height = 30
        worksheet.row_dimensions[2].height = 20
        
        for col_idx, col_name in enumerate(ordered_columns, 1):
            col_letter = get_column_letter(col_idx)
            max_length = max(
                len(str(col_name)),
                df_ordered[col_name].astype(str).apply(len).max() if len(df_ordered) > 0 else 0
            )
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[col_letter].width = adjusted_width
        
        worksheet.freeze_panes = 'A3'
    
    output.seek(0)
    return output.getvalue()


def to_csv_bytes(df: pd.DataFrame, original_columns_mapping: Optional[Dict[str, str]] = None) -> bytes:
    """Convert DataFrame ke CSV bytes."""
    if original_columns_mapping:
        df = restore_original_columns(df, original_columns_mapping)
    
    return df.to_csv(index=False).encode('utf-8')
